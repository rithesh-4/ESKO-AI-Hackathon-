"""Generates realistic synthetic sample files (no Esko mail needed)."""
from fpdf import FPDF
import openpyxl
from PIL import Image, ImageDraw

# 1. spec PDF
pdf = FPDF()
pdf.add_page()
pdf.set_font("Helvetica", "B", 14)
pdf.cell(0, 10, "Product Spec - Mango Oats Bar (DRAFT)")
pdf.ln()
pdf.set_font("Helvetica", "", 11)
pdf.multi_cell(0, 6,
"Product Name: Mango Oats Bar\nCategory: Snacks - Cereal Bar\nPack format: pillow pouch\n"
"Pack size: 45g (see nutrition sheet - CHECK)\nIngredients: Oats (42%), mango pulp (18%), honey, rice crisps, milk solids, cashew\n"
"Claims: No added refined sugar; Source of fibre\nBranding: keep existing SunnyBite logo unchanged? TBC with brand manager\n"
"Target market: India (pan-India, English + Hindi back panel)\nMandatory: FSSAI logo + Lic No, veg mark, MRP, Mfg/Expiry, customer care, Mfd by address\n"
"Note: allergen declaration missing in this draft. Tagline: direction TBD.")
pdf.output("samples/product_spec.pdf")

# 2. nutrition xlsx (intentional 50g conflict)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Nutrition"
for row in [
    ["Product", "Mango Oats Bar"],
    ["Serving / Net wt", "50g (conflict: spec says 45g)"],
    ["Nutrient", "Per 100g", "Per serve"],
    ["Energy", "410 kcal", "205 kcal"],
    ["Protein", "9g", "4.5g"],
    ["Carbohydrate", "68g", "34g"],
    ["Total Sugars", "24g", "12g"],
    ["Fat", "11g", "5.5g"],
    ["Fibre", "7g", "3.5g"],
]:
    ws.append(row)
ws2 = wb.create_sheet("PackCopy")
for row in [
    ["Field", "Value"],
    ["Mfd by", "SunnyBite Foods, Plot 12 MIDC Pune"],
    ["Customer care", "1800-xxx-xxx, care@sunnybite.in"],
    ["MRP", "Rs 35 (inclusive of all taxes)"],
]:
    ws2.append(row)
wb.save("samples/nutrition.xlsx")

# 3. mock old pack image
img = Image.new("RGB", (600, 800), (255, 243, 205))
d = ImageDraw.Draw(img)
d.rectangle([40, 40, 560, 760], outline=(180, 120, 40), width=6)
d.text((80, 120), "SunnyBite", fill=(160, 60, 10))
d.text((80, 200), "MANGO OATS BAR", fill=(60, 60, 60))
d.text((80, 260), "45g?", fill=(180, 0, 0))
d.text((80, 340), "Old pack - front", fill=(100, 100, 100))
d.text((80, 420), "No allergen panel visible", fill=(150, 0, 0))
img.save("samples/old_pack.jpg")
print("samples written")
