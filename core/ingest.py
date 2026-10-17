"""File ingestion: PDF / DOCX / XLSX / images -> list of text chunks with provenance."""
from __future__ import annotations
import io
from dataclasses import dataclass


@dataclass
class Chunk:
    file: str
    loc: str
    text: str


def extract_pdf(data: bytes, filename: str) -> list[Chunk]:
    import pymupdf  # PyMuPDF
    doc = pymupdf.open(stream=data, filetype="pdf")
    out: list[Chunk] = []
    for i, page in enumerate(doc, start=1):
        t = page.get_text().strip()
        if t:
            out.append(Chunk(filename, f"p{i}", t))
    return out


def extract_docx(data: bytes, filename: str) -> list[Chunk]:
    import docx
    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    tables = []
    for ti, t in enumerate(d.tables):
        for r in t.rows:
            tables.append(" | ".join(c.text.strip() for c in r.cells))
    full = text + ("\n" + "\n".join(tables) if tables else "")
    return [Chunk(filename, "doc", full)] if full.strip() else []


def extract_xlsx(data: bytes, filename: str) -> list[Chunk]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out: list[Chunk] = []
    for ws in wb.worksheets:
        lines = []
        for row in ws.iter_rows(values_only=True):
            vals = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if vals:
                lines.append(" | ".join(vals))
        if lines:
            out.append(Chunk(filename, f"Sheet:{ws.title}", "\n".join(lines)))
    return out


def extract_txt(data: bytes, filename: str) -> list[Chunk]:
    try:
        t = data.decode("utf-8", errors="ignore").strip()
    except Exception:
        t = ""
    return [Chunk(filename, "txt", t)] if t else []


def extract_image(data: bytes, filename: str) -> list[Chunk]:
    # Offline-safe: record image presence; vision LLM enriches later in llm.py.
    # Try local OCR if tesseract binary exists, else return placeholder.
    text = ""
    try:
        from PIL import Image
        import pytesseract
        img = Image.open(io.BytesIO(data))
        text = pytesseract.image_to_string(img).strip()
    except Exception:
        text = ""
    note = text if text else "[image: old pack design — visual review needed]"
    return [Chunk(filename, "image", note)]


def ingest(filename: str, data: bytes) -> list[Chunk]:
    fn = filename.lower()
    if fn.endswith(".pdf"):
        return extract_pdf(data, filename)
    if fn.endswith(".docx"):
        return extract_docx(data, filename)
    if fn.endswith(".xlsx") or fn.endswith(".xls"):
        return extract_xlsx(data, filename)
    if fn.endswith(".txt") or fn.endswith(".md"):
        return extract_txt(data, filename)
    if fn.endswith((".png", ".jpg", ".jpeg", ".webp")):
        return extract_image(data, filename)
    # fallback: try text
    return extract_txt(data, filename)
